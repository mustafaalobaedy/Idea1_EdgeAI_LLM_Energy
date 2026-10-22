"""Build the PeerJ-template Word manuscript from paper/sections/*.md + references.bib.

Sections are concatenated in file-name order (files whose name contains "outline" are skipped).

Draft syntax
  TITLE: / ABSTRACT: / KEYWORDS:  front matter (00_front.md)
  "# ", "## ", "### "              heading levels 1-3
  "- BULLET: text"                 bullet item
  "HIGHLIGHT: text"                paragraph highlighted in yellow (AI-use statement)
  {{FIG:path|caption|label}}       picture (PNG) with caption below
  {{DIAGRAM:name|caption|label}}   editable DrawingML diagram from diagrams.py
  {{TABLE:path|caption|label}}     pipe-table (.md) or .csv file with caption above
  {{EQ:name}}                      numbered Word (OMML) equation from equations.py
  {{ALGO:caption|label}} ... {{/ALGO}}   pseudocode box (lines; two leading spaces = one indent level)
  [fig:x] [tab:x] [eq:x] [alg:x]   cross-references -> Figure n / Table n / (n) / Algorithm n
  [@key] or [@a; @b]               author-year citations (PeerJ); "Lastname et al. [@key]" becomes narrative
  *italic*, **bold**, `code`       inline markup
The build fails if a label is referenced but undefined, defined but never referenced, or a citation key is missing.

Usage: python build_paper.py [--version V02]   (writes paper/paper.docx; with --version also a named copy + PDF)
"""
import csv
import glob
import os
import re
import subprocess
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
TEMPLATE = Path(os.environ.get('PAPER_TEMPLATE', ROOT.parent / 'PeerJ-research-manuscript-template.docx'))
OUT = HERE / 'paper.docx'
FONT = 'Times New Roman'
BODY_PT, TABLE_PT, CAPTION_PT = 12, 10, 11
TEXT_WIDTH_IN = 6.5

sys.path.insert(0, str(HERE))
from equations import EQS  # noqa: E402
from diagrams import DIAGRAMS  # noqa: E402


# ---------------------------------------------------------------- bibliography
def parse_bib(text):
    entries, i = {}, 0
    while True:
        m = re.compile(r'@(\w+)\s*\{\s*([^,\s]+)\s*,').search(text, i)
        if not m:
            return entries
        typ, key, j, depth = m.group(1).lower(), m.group(2), m.end(), 1
        start = j
        while depth:
            depth += {'{': 1, '}': -1}.get(text[j], 0)
            j += 1
        body, fields = text[start:j - 1], {}
        for fm in re.finditer(r'(\w+)\s*=\s*\{', body):
            k, d = fm.end(), 1
            while d:
                d += {'{': 1, '}': -1}.get(body[k], 0)
                k += 1
            fields[fm.group(1).lower()] = body[fm.end():k - 1]
        entries[key] = dict(type=typ, **fields)
        i = j


_ACC = {"\\'": {'a': 'á', 'e': 'é', 'i': 'í', 'o': 'ó', 'u': 'ú', 'c': 'ć', 'n': 'ń', 'y': 'ý'},
        '\\"': {'a': 'ä', 'e': 'ë', 'i': 'ï', 'o': 'ö', 'u': 'ü'},
        '\\`': {'a': 'à', 'e': 'è', 'o': 'ò', 'u': 'ù'},
        '\\^': {'a': 'â', 'e': 'ê', 'o': 'ô', 'i': 'î'},
        '\\~': {'a': 'ã', 'n': 'ñ', 'o': 'õ'},
        '\\c': {'c': 'ç'}, '\\v': {'c': 'č', 's': 'š', 'z': 'ž', 'r': 'ř'}, '\\u': {'g': 'ğ'}}


def clean(s):
    s = s.replace('\\aa', 'å').replace('\\o', 'ø').replace('\\ss', 'ß').replace('\\&', '&').replace('\\_', '_')
    s = s.replace('$\\tau$', 'τ').replace('--', '–').replace('``', '“').replace("''", '”')
    for acc, table in _ACC.items():
        for ch, out in table.items():
            s = s.replace('{' + acc + ch + '}', out).replace('{' + acc + '{' + ch + '}}', out).replace(acc + ch, out)
    return re.sub(r'[{}]', '', s).strip()


def split_authors(s):
    names = [a.strip() for a in re.split(r'\s+and\s+', s) if a.strip()]
    out, et_al = [], False
    for n in names:
        if n == 'others':
            et_al = True
            continue
        if n.startswith('{') and n.endswith('}'):
            out.append((clean(n), ''))  # corporate author
        elif ',' in n:
            last, first = [x.strip() for x in n.split(',', 1)]
            out.append((clean(last), clean(first)))
        else:
            parts = clean(n).split()
            out.append((parts[-1], ' '.join(parts[:-1])))
    return out, et_al


def initials(first):
    return ''.join(p[0] for p in re.split(r'[\s.\-]+', first) if p and p[0].isupper())


def cite_name(e):
    """Author part of an in-text citation: Smith / Smith & Jones / Smith et al."""
    au, et_al = split_authors(e['author'])
    if len(au) == 1 and not et_al:
        return au[0][0]
    if len(au) == 2 and not et_al:
        return f'{au[0][0]} & {au[1][0]}'
    return f'{au[0][0]} et al.'


def peerj_ref(e, year_label):
    """PeerJ reference-list entry as (text, italic) segments."""
    au, et_al = split_authors(e['author'])
    names = ', '.join(f'{l} {initials(f)}'.strip() for l, f in au) + (' et al.' if et_al else '')
    t = clean(e['title']).rstrip('.')
    doi = f' DOI {e["doi"]}.' if 'doi' in e else ''
    url = e.get('url', '')
    note = clean(e.get('note', ''))
    acc = f' (accessed {note.split("Accessed ")[-1]})' if 'Accessed' in note else ''
    head = [(f'{names}. {year_label}. {t}. ', False)]
    typ = e['type']
    if typ == 'article':
        vol = e.get('volume', '')
        no = f'({e["number"]})' if 'number' in e else ''
        pg = clean(e.get('pages', ''))
        tail = f' {vol}{no}' + (f':{pg}' if pg else '') if vol else (f' {pg}' if pg else '')
        return head + [(clean(e['journal']), True), (tail + (doi if doi else '.'), False)]
    if typ in ('inproceedings', 'incollection'):
        pg = clean(e.get('pages', ''))
        pub = clean(e.get('publisher', ''))
        tail = ''.join([f', {pg}' if pg else '', f'. {pub}' if pub else '', doi if doi else '.'])
        return head + [('In: ', False), (clean(e['booktitle']), True), (tail, False)]
    if 'eprint' in e:
        return head + [(f'arXiv preprint arXiv:{e["eprint"]}.' + (doi if doi else ''), False)]
    hp = clean(e.get('howpublished', ''))
    seg = (f'{hp}. ' if hp else '') + (f'Available at {url}{acc}.' if url else '') + (doi if doi and not url else '')
    return head + [(seg, False)]


# ---------------------------------------------------------------- inline markup
def inline(text):
    """*italic*, **bold**, `code` -> (text, italic, bold, code) segments."""
    segs = []
    for j, part in enumerate(re.split(r'\*\*([^*]+)\*\*', text)):
        bold = j % 2 == 1
        for i, piece in enumerate(re.split(r'\*([^*]+)\*', part)):
            it = i % 2 == 1
            for k, x in enumerate(re.split(r'`([^`]+)`', piece)):
                if x:
                    segs.append((x, it, bold, k % 2 == 1))
    return segs


class Builder:
    def __init__(self):
        self.d = docx.Document(TEMPLATE)
        body = self.d.element.body
        for el in list(body.iterchildren()):
            if el.tag != qn('w:sectPr'):
                body.remove(el)
        self.pic_id = 100

    # ------------------------------------------------------------ paragraphs
    def para(self, segs, size=BODY_PT, align=None, before=0, after=8, bold=False, italic=False, keep_next=False,
             left_indent=None, hanging=None, line=None, highlight=False):
        p = self.d.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        if line:
            pf.line_spacing = line
        if align is not None:
            p.alignment = align
        if keep_next:
            pf.keep_with_next = True
        if left_indent is not None:
            pf.left_indent = Inches(left_indent)
        if hanging is not None:
            pf.first_line_indent = Inches(-hanging)
        for seg in segs:
            if isinstance(seg, str):
                seg = (seg, False, False, False)
            text, it, bd, code = (list(seg) + [False, False, False])[:4]
            r = p.add_run(text)
            r.font.name = 'Consolas' if code else FONT
            r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
            r.font.size = Pt(size - 1 if code else size)
            r.italic = it or italic
            r.bold = bd or bold
            if highlight:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return p

    def heading(self, text, level):
        size = {1: 14, 2: 12, 3: 12, 4: 12}[level]
        self.para([(text, level >= 3, level <= 3, False)], size=size, before={1: 18, 2: 12, 3: 8, 4: 6}[level], after=6, keep_next=True)

    def bullet(self, segs):
        self.para([('•\t', False, False, False)] + segs, left_indent=0.3, hanging=0.2, after=4)

    def caption(self, prefix, text):
        self.para([(prefix, False, True, False)] + inline(text), size=CAPTION_PT, after=10, before=4)

    # ------------------------------------------------------------ equations
    def equation(self, name, number):
        p = self.d.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(6), Pt(8)
        pf.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN / 2), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run('\t')
        p._p.append(parse_xml(EQS[name]))
        r = p.add_run('\t' + number)
        r.font.name, r.font.size = FONT, Pt(BODY_PT)

    # ------------------------------------------------------------ figures / diagrams
    def figure(self, path, caption, number):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(8)
        p.add_run().add_picture(str(ROOT / path), width=Inches(min(TEXT_WIDTH_IN, 6.0)))
        self.caption(f'Figure {number}: ', caption)

    def diagram(self, name, caption, number):
        self.pic_id += 1
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(8)
        p._p.append(parse_xml(DIAGRAMS[name]().xml(self.pic_id, f'Figure {number}')))
        self.caption(f'Figure {number}: ', caption)

    # ------------------------------------------------------------ tables
    @staticmethod
    def _borders(tbl, header_rows=1, caption_row=False):
        tblPr = tbl._tbl.tblPr
        b = OxmlElement('w:tblBorders')
        for side in ('top', 'bottom', 'left', 'right', 'insideV', 'insideH'):
            e = OxmlElement(f'w:{side}')
            if side == 'bottom':
                e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '8'); e.set(qn('w:color'), '000000')
            else:
                e.set(qn('w:val'), 'nil')
            b.append(e)
        tblPr.append(b)
        rule_rows = [header_rows - 1] + ([0] if caption_row else [])
        for ri in rule_rows:
            for c in tbl.rows[ri].cells:
                tcPr = c._tc.get_or_add_tcPr()
                tb = OxmlElement('w:tcBorders')
                e = OxmlElement('w:bottom')
                e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '8' if ri == 0 and caption_row else '6'); e.set(qn('w:color'), '000000')
                tb.append(e)
                tcPr.append(tb)

    def table(self, path, caption, number):
        full = ROOT / path
        if full.suffix == '.csv':
            rows = [r for r in csv.reader(open(full))]
        else:
            rows = [[c.strip() for c in r.strip().strip('|').split('|')] for r in full.read_text().splitlines()
                    if r.strip() and not re.match(r'^\|?\s*:?-', r.strip())]
        ncol = max(len(r) for r in rows)
        rows = [r + [''] * (ncol - len(r)) for r in rows]
        size = TABLE_PT if ncol <= 6 else 9 if ncol <= 8 else 8
        # column widths: every column gets at least its longest unbreakable word; the rest is shared by content length
        char_in = 0.0088 * size  # rough average character width in inches for Times at `size` pt
        min_w = [max(len(w) for r in rows for w in (r[c].split() or [''])) * char_in + 0.16 for c in range(ncol)]
        want = [max(len(r[c]) for r in rows) * char_in + 0.16 for c in range(ncol)]
        want = [min(w, 2.6) for w in want]
        free = TEXT_WIDTH_IN - sum(min_w)
        extra = [max(w - m, 0) for w, m in zip(want, min_w)]
        widths = [m + (free * e / sum(extra) if free > 0 and sum(extra) > 0 else 0) for m, e in zip(min_w, extra)]
        if sum(widths) > TEXT_WIDTH_IN:
            widths = [w * TEXT_WIDTH_IN / sum(widths) for w in widths]
        widths = [Inches(w) for w in widths]
        t = self.d.add_table(rows=len(rows) + 1, cols=ncol)
        t.autofit = False
        # caption row (merged across all columns) so that the caption can never be separated from its table
        cap_cell = t.cell(0, 0).merge(t.cell(0, ncol - 1))
        p = cap_cell.paragraphs[0]
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(4)
        p.paragraph_format.keep_with_next = True
        for text, it, bd, code in [(f'Table {number}: ', False, True, False)] + inline(caption):
            r = p.add_run(text)
            r.font.name, r.font.size, r.italic, r.bold = FONT, Pt(CAPTION_PT), it, bd
        long_table = len(rows) > 18
        for ri, row in enumerate(rows):
            for ci, cell in enumerate(row):
                c = t.cell(ri + 1, ci)
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                if ri == 0 or (not long_table and ri < len(rows) - 1):
                    p.paragraph_format.keep_with_next = True
                for text, it, bd, code in inline(cell):
                    r = p.add_run(text)
                    r.font.name, r.font.size = FONT, Pt(size)
                    r.italic, r.bold = it, bd or ri == 0
                if ci >= 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for ri in range(1, len(rows) + 1):
            for ci in range(ncol):
                t.cell(ri, ci).width = widths[ci]
        cap_cell.width = Inches(sum(w.inches for w in widths))
        for tr in t._tbl.findall(qn('w:tr')):
            trPr = tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:cantSplit'))
        self._borders(t, header_rows=2, caption_row=True)
        tblPr = t._tbl.tblPr
        mar = OxmlElement('w:tblCellMar')
        for side in ('left', 'right'):
            m_ = OxmlElement(f'w:{side}')
            m_.set(qn('w:w'), '50'); m_.set(qn('w:type'), 'dxa')
            mar.append(m_)
        tblPr.append(mar)
        for tag, attrs in (('w:tblW', {'w:w': str(sum(w.twips for w in widths)), 'w:type': 'dxa'}),
                           ('w:tblLayout', {'w:type': 'fixed'}), ('w:jc', {'w:val': 'center'})):
            for old_el in tblPr.findall(qn(tag)):
                tblPr.remove(old_el)
            el = OxmlElement(tag)
            for k_, v_ in attrs.items():
                el.set(qn(k_), v_)
            tblPr.append(el)
        for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(w.twips))
        self.para([], after=6)

    # ------------------------------------------------------------ algorithms
    def algorithm(self, caption, lines, number):
        t = self.d.add_table(rows=2, cols=1)
        head = t.cell(0, 0).paragraphs[0]
        head.paragraph_format.space_before = Pt(2)
        head.paragraph_format.space_after = Pt(2)
        head.paragraph_format.keep_with_next = True
        for text, it, bd, code in [(f'Algorithm {number}: ', False, True, False)] + inline(caption):
            r = head.add_run(text)
            r.font.name, r.font.size, r.italic, r.bold = FONT, Pt(TABLE_PT), it, bd
        cell = t.cell(1, 0)
        first, k = True, 0
        for ln in lines:
            if not ln.strip():
                continue
            k += 1
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            indent = (len(ln) - len(ln.lstrip(' '))) // 2
            p.paragraph_format.left_indent = Inches(0.15 * indent + 0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            for text, it, bd, code in [(f'{k}: ', False, False, False)] + inline(ln.strip()):
                r = p.add_run(text)
                r.font.name, r.font.size, r.italic, r.bold = FONT, Pt(TABLE_PT), it, bd
        self._borders(t)
        tblPr = t._tbl.tblPr
        for tag, attrs in (('w:tblW', {'w:w': str(int(TEXT_WIDTH_IN * 1440)), 'w:type': 'dxa'}), ('w:tblLayout', {'w:type': 'fixed'})):
            el = OxmlElement(tag)
            for k_, v_ in attrs.items():
                el.set(qn(k_), v_)
            tblPr.append(el)
        for gc in t._tbl.tblGrid.findall(qn('w:gridCol')):
            gc.set(qn('w:w'), str(int(TEXT_WIDTH_IN * 1440)))
        for row in t.rows:
            row.cells[0].width = Inches(TEXT_WIDTH_IN)
        self.para([], after=6)


# ---------------------------------------------------------------- main
def main():
    version = sys.argv[sys.argv.index('--version') + 1] if '--version' in sys.argv else None
    bib = parse_bib((HERE / 'references.bib').read_text())
    files = sorted(f for f in glob.glob(str(HERE / 'sections' / '*.md')) if 'outline' not in os.path.basename(f))
    draft = '\n\n'.join(Path(f).read_text() for f in files)

    # ---- citations: collect keys, assign year labels (a/b for same name+year), render author-year
    used_keys = []
    for m in re.finditer(r'\[(@[^\]]+)\]', draft):
        for key in [k.strip().lstrip('@') for k in m.group(1).split(';')]:
            if key not in bib:
                raise KeyError(f'citation key not in references.bib: {key}')
            if key not in used_keys:
                used_keys.append(key)
    ordered = sorted(used_keys, key=lambda k: (split_authors(bib[k]['author'])[0][0][0].lower(), bib[k]['year'], clean(bib[k]['title']).lower()))
    year_label = {}
    for k in ordered:
        same = [j for j in ordered if cite_name(bib[j]) == cite_name(bib[k]) and bib[j]['year'] == bib[k]['year']]
        year_label[k] = bib[k]['year'] + (chr(ord('a') + same.index(k)) if len(same) > 1 else '')

    def narrative_or_paren(m):
        prefix, keys = m.group(1) or '', [k.strip().lstrip('@') for k in m.group(2).split(';')]
        if prefix and len(keys) == 1:
            k = keys[0]
            first_last = split_authors(bib[k]['author'])[0][0][0]
            pm = re.match(r"([A-Z][\w'’\-]+)( (?:and|&) [A-Z][\w'’\-]+)?( et al\.)?$", prefix.strip())
            if pm and pm.group(1) == first_last.split()[-1]:
                name = cite_name(bib[k])
                return f'{name} ({year_label[k]})'
        return (prefix + '(' + '; '.join(f'{cite_name(bib[k])}, {year_label[k]}' for k in keys) + ')')
    draft = re.sub(r"((?:[A-Z][\w'’\-]+(?: (?:and|&) [A-Z][\w'’\-]+)?(?: et al\.)? )?)\[(@[^\]]+)\]", narrative_or_paren, draft)

    # ---- numbering of figures, tables, equations and algorithms in order of definition
    labels, counters = {}, {'fig': 0, 'tab': 0, 'eq': 0, 'alg': 0}
    for m in re.finditer(r'\{\{(FIG|DIAGRAM|TABLE|EQ|ALGO):([^}]*)\}\}', draft):
        kind = {'FIG': 'fig', 'DIAGRAM': 'fig', 'TABLE': 'tab', 'EQ': 'eq', 'ALGO': 'alg'}[m.group(1)]
        lab = m.group(2) if kind == 'eq' else m.group(2).split('|')[-1].strip()
        if f'{kind}:{lab}' in labels:
            raise KeyError(f'duplicate label {kind}:{lab}')
        counters[kind] += 1
        n = counters[kind]
        labels[f'{kind}:{lab}'] = {'fig': f'Figure {n}', 'tab': f'Table {n}', 'eq': f'({n})', 'alg': f'Algorithm {n}'}[kind]
    body_text = re.sub(r'\{\{[^}]*\}\}', '', draft)
    used = set(re.findall(r'\[((?:fig|tab|eq|alg):[\w-]+)\]', body_text))
    if used - set(labels):
        raise KeyError(f'referenced but undefined: {sorted(used - set(labels))}')
    if set(labels) - used:
        raise KeyError(f'defined but never referenced in the text: {sorted(set(labels) - used)}')
    draft = re.sub(r'\[((?:fig|tab|eq|alg):[\w-]+)\]', lambda m: labels[m.group(1)], draft)
    if re.search(r'\[TO (WRITE|FINALI[SZ]E)', draft):
        print('WARNING: draft still contains [TO WRITE] placeholders')

    B = Builder()
    algo = None
    long_sentences = 0
    for line in draft.splitlines():
        s = line.strip()
        if algo is not None:
            if s.startswith('{{/ALGO}}'):
                B.algorithm(algo[0], algo[1], algo[2])
                algo = None
            else:
                algo[1].append(line.rstrip())
            continue
        if not s:
            continue
        if s.startswith('{{') and not s.endswith('}}'):
            raise ValueError(f'text after a {{...}} directive would be lost: {s[:80]}')
        if not s.startswith(('{{', 'TITLE:', 'AUTHOR:', '#')):
            plain = re.sub(r'\([^)]*\d{4}[a-z]?\)|e\.g\.|i\.e\.|et al\.|vs\.', '', s)
            for sent in re.split(r'(?<=[.?!])\s+(?=[A-Z(])', plain):
                if len(sent.split()) > 30:
                    long_sentences += 1
                    print(f'LONG SENTENCE ({len(sent.split())} words): {sent[:90]}...')
        if s.startswith('{{ALGO:'):
            cap, lab = s[7:-2].split('|')[:2]
            algo = (cap, [], labels[f'alg:{lab.strip()}'].split()[-1])
        elif s.startswith('{{EQ:'):
            B.equation(s[5:-2], labels[f'eq:{s[5:-2]}'])
        elif s.startswith('{{FIG:'):
            path, cap, lab = [x.strip() for x in s[6:-2].split('|')[:3]]
            B.figure(path, cap, labels[f'fig:{lab}'].split()[-1])
        elif s.startswith('{{DIAGRAM:'):
            name, cap, lab = [x.strip() for x in s[10:-2].split('|')[:3]]
            B.diagram(name, cap, labels[f'fig:{lab}'].split()[-1])
        elif s.startswith('{{TABLE:'):
            path, cap, lab = [x.strip() for x in s[8:-2].split('|')[:3]]
            B.table(path, cap, labels[f'tab:{lab}'].split()[-1])
        elif s.startswith('TITLE:'):
            B.para([(s[6:].strip(), False, True, False)], size=18, after=18)
        elif s.startswith('AUTHOR:'):
            name, aff, addr, email = [x.strip() for x in s[7:].split('|')]
            B.para([(name, False, False, False), ('1', False, False, False)], after=6)
            B.d.paragraphs[-1].runs[-1].font.superscript = True
            B.para([('1 ', False, False, False), (aff, False, False, False)], after=12)
            B.para([('Corresponding Author:', False, True, False)], after=0)
            B.para([(name, False, False, False)], after=0)
            B.para([(addr, False, False, False)], after=0)
            B.para([('Email address: ', False, False, False), (email, False, False, False)], after=12)
        elif s.startswith('ABSTRACT:'):
            B.heading('Abstract', 1)
            B.para(inline(s[9:].strip()))
        elif s.startswith('KEYWORDS:'):
            B.para([('Keywords: ', False, True, False)] + inline(s[9:].strip()))
        elif s.startswith('#### '):
            B.heading(s[5:], 4)
        elif s.startswith('### '):
            B.heading(s[4:], 3)
        elif s.startswith('## '):
            B.heading(s[3:], 2)
        elif s.startswith('# '):
            B.heading(s[2:], 1)
        elif s.startswith('- BULLET:'):
            B.bullet(inline(s[9:].strip()))
        elif s.startswith('HIGHLIGHT:'):
            B.para(inline(s[10:].strip()), highlight=True)
        else:
            B.para(inline(s))

    B.heading('References', 1)
    for key in ordered:
        segs = [(t, it, False, False) for t, it in peerj_ref(bib[key], year_label[key])]
        B.para(segs, size=11, after=6, left_indent=0.3, hanging=0.3)

    B.d.save(OUT)
    unused = sorted(set(bib) - set(ordered))
    print(f'wrote {OUT}  ({len(ordered)} references cited; {len(unused)} bib entries unused; {long_sentences} long sentences)')
    if version:
        title = re.search(r'^TITLE:\s*(.+)$', draft, re.M).group(1)
        words = re.sub(r'[^\w\s]', '', title).split()[:5]
        name = ' '.join(words) + f' {version}.docx'
        (HERE / name).write_bytes(OUT.read_bytes())
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', str(HERE), str(HERE / name)],
                       capture_output=True, text=True, timeout=300)
        print(f'versioned copy: {name} (+ PDF)')


if __name__ == '__main__':
    main()
